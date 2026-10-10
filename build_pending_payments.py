"""Builds Pending_Payments.xlsx, a replica of Pending_Payments_Final.numbers.

Data rows live in A11:G106; Days Open (H) and Status (I) are formulas.
Edit the DATA list below (or the xlsx directly) and keep using the xlsx going forward.
"""
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

DATA = [
    ("Fiorentino Thomas", "File# 1699.0041", "Jeffrey Magdalen", "Gerber Ciano Kelly Brady LLP", date(2026, 8, 19), 3600),
    ("Ellis Brenda", "File# 3693.00246", "Susan Miller", "RPG", date(2026, 8, 26), 2625),
    ("Fraser Beverly", "File# 46.12500", "Nick Albanese", "RPG", date(2026, 8, 30), 3600),
    ("LeRose Sondra", "File# P6529.2428", "Nick Albanese", "RPG", date(2026, 8, 30), 3300),
    ("Stamm #2", "File# 3693.00246", "Anna Borea", "Wilson Elser", date(2026, 9, 11), 1650),
    ("Bambach Trial Prep", "File: PB-1044", "Christian Fumoso", "Fumoso Christesen", date(2026, 9, 16), 1100),
    ("Wakefield Angela", "File# 14583.00120", "Susan Formosa", "Wilson Elser", date(2026, 9, 25), 5775),
    ("Biesty Kevin", "File: 21975-00108", "Jillian Oliff", "Wilson Elser", date(2026, 9, 29), 4500),
    ("Pettway", "BPN File No: 0102-087", "Brian DiPentima", "BPN LLP", date(2026, 9, 30), 5100),
    ("Power Patricia", "File: 33-225", "Eric Doman", "Wagner, Doman, Leto & DiLeo P.C.", date(2026, 10, 6), 1500),
]
FIRST, LAST = 11, 106
BUCKETS = ["Under 30 Days", "31–60 Days", "61–90 Days", "91–120 Days", "121–150 Days", "151+ Days"]

wb = Workbook()
ws = wb.active
ws.title = "Pending Payments"
NAVY = "1F3A5F"
bold = Font(bold=True)
hdr_fill = PatternFill("solid", fgColor=NAVY)
hdr_font = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="BBBBBB")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
money = '"$"#,##0.00'

ws["A2"] = "Pending Payments"; ws["A2"].font = Font(bold=True, size=20, color=NAVY)
ws["H2"] = "As of"; ws["H2"].font = bold; ws["H2"].alignment = Alignment(horizontal="right")
ws["I2"] = "=TODAY()"; ws["I2"].number_format = "mmm d, yyyy"
ws["A3"] = "Enter Y in Paid when an invoice is settled. Aging and totals update automatically each day."
ws["A3"].font = Font(italic=True, color="666666")

for col, (lab, sub) in zip("ACEG", [("TOTAL OUTSTANDING", "Unpaid balance"), ("OPEN INVOICES", "Active items"),
                                    ("AVERAGE INVOICE", "Per open invoice"), ("AVERAGE DAYS OPEN", "Dated open invoices")]):
    ws[f"{col}5"] = lab; ws[f"{col}5"].font = Font(bold=True, size=9, color="666666")
    ws[f"{col}7"] = sub; ws[f"{col}7"].font = Font(size=9, color="888888")
    ws[f"{col}6"].font = Font(bold=True, size=16, color=NAVY)
ws["A6"] = "=SUM($M$4:$M$10)"; ws["A6"].number_format = money
ws["C6"] = "=SUM($L$4:$L$10)"
ws["E6"] = "=IF($C$6=0,0,$A$6/$C$6)"; ws["E6"].number_format = money
ws["G6"] = f"=IF(COUNT(H{FIRST}:H{LAST})=0,0,SUM(H{FIRST}:H{LAST})/COUNT(H{FIRST}:H{LAST}))"; ws["G6"].number_format = "0.0"
for c in "ACEG":
    ws[f"{c}6"].alignment = Alignment(horizontal="left")

ws["K2"] = "AGING SUMMARY"; ws["K2"].font = Font(bold=True, color=NAVY)
for c, t in zip("KLMN", ["Age Bucket", "Invoices", "Outstanding", "% of Total"]):
    ws[f"{c}3"] = t; ws[f"{c}3"].font = hdr_font; ws[f"{c}3"].fill = hdr_fill
for i, b in enumerate(BUCKETS):
    r = 4 + i
    ws[f"K{r}"] = b
    ws[f"L{r}"] = f"=COUNTIF($I${FIRST}:$I${LAST},K{r})"
    ws[f"M{r}"] = f"=SUMIF($I${FIRST}:$I${LAST},K{r},$F${FIRST}:$F${LAST})"; ws[f"M{r}"].number_format = money
    ws[f"N{r}"] = f"=IF($A$6=0,0,M{r}/$A$6)"; ws[f"N{r}"].number_format = "0.0%"
    for c in "KLMN": ws[f"{c}{r}"].border = box
ws["K10"] = "Date Needed"
ws["L10"] = f"=COUNTIF($I${FIRST}:$I${LAST},K10)"
ws["M10"] = f"=SUMIF($I${FIRST}:$I${LAST},K10,$F${FIRST}:$F${LAST})"; ws["M10"].number_format = money
ws["N10"] = "=IF($A$6=0,0,M10/$A$6)"; ws["N10"].number_format = "0.0%"
for c in "KLMN": ws[f"{c}10"].border = box

for c, t in zip("ABCDEFGHI", ["Case Name", "File Number", "Attorney Name", "Law Firm", "Invoice Sent", "Amount", "Paid", "Days Open", "Status"]):
    ws[f"{c}10"] = t; ws[f"{c}10"].font = hdr_font; ws[f"{c}10"].fill = hdr_fill
    ws[f"{c}10"].alignment = Alignment(horizontal="center")
ws["K10"].font = Font(); # summary row label keeps plain style

for r in range(FIRST, LAST + 1):
    i = r - FIRST
    if i < len(DATA):
        n, f, a, l, d, amt = DATA[i]
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"F{r}"], ws[f"G{r}"] = n, f, a, l, d, amt, "N"
    ws[f"E{r}"].number_format = "mm/dd/yyyy"
    ws[f"F{r}"].number_format = money
    ws[f"G{r}"].alignment = Alignment(horizontal="center")
    ws[f"H{r}"] = f'=IF(A{r}="","",IF(UPPER(G{r})="Y","",IF(E{r}="","",TODAY()-E{r})))'
    ws[f"I{r}"] = (f'=IF(A{r}="","",IF(UPPER(G{r})="Y","Paid",IF(E{r}="","Date Needed",'
                   f'IF(H{r}<=30,"Under 30 Days",IF(H{r}<=60,"31–60 Days",IF(H{r}<=90,"61–90 Days",'
                   f'IF(H{r}<=120,"91–120 Days",IF(H{r}<=150,"121–150 Days","151+ Days"))))))))')
    ws[f"H{r}"].number_format = "0"
    for c in "ABCDEFGHI": ws[f"{c}{r}"].border = box

# paid rows grey out; older invoices tint
rng = f"A{FIRST}:I{LAST}"
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'UPPER($G{FIRST})="Y"'], font=Font(color="999999"), fill=PatternFill("solid", bgColor="EEEEEE")))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($A{FIRST}<>"",ISNUMBER($H{FIRST}),$H{FIRST}>60)'], fill=PatternFill("solid", bgColor="F8D7DA")))
ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($A{FIRST}<>"",ISNUMBER($H{FIRST}),$H{FIRST}>30)'], fill=PatternFill("solid", bgColor="FFF3CD")))

dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"G{FIRST}:G{LAST}")

ws["K12"] = "HOW TO USE"; ws["K12"].font = bold
ws["K13"] = "Enter Y in Paid when the invoice is settled."
ws["K14"] = "Formulas, aging, and totals update automatically."
ws["K16"] = "Total Outstanding"; ws["M16"] = "=$A$6"; ws["M16"].number_format = money
ws["K18"] = "Average Invoice"; ws["M18"] = "=$E$6"; ws["M18"].number_format = money
ws["K19"] = "Average Days Open"; ws["M19"] = "=$G$6"; ws["M19"].number_format = "0.0"
for c in ("K16", "K18", "K19"): ws[c].font = bold

for col, w in zip("ABCDEFGHIJKLMN", [22, 22, 20, 34, 13, 13, 7, 11, 15, 3, 18, 10, 14, 11]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A11"
wb.save("Pending_Payments.xlsx")
